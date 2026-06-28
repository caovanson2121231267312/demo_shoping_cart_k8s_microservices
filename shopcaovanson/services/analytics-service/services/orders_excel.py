import io
from datetime import datetime, timedelta, timezone
from typing import Any, Callable, Iterator

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from db import order_db

ProgressCallback = Callable[[int, str], None]

STATUS_LABELS = {
    "pending": "Chờ xử lý",
    "confirmed": "Đã xác nhận",
    "shipping": "Đang giao",
    "delivered": "Đã giao",
    "cancelled": "Đã hủy",
}

FETCH_BATCH = 5000

VND_FMT = '#,##0" ₫"'
INT_FMT = "#,##0"
DATETIME_FMT = "dd/mm/yyyy hh:mm"
DATE_FMT = "dd/mm/yyyy"

TITLE_FONT = Font(bold=True, size=16, color="1B5E20")
SUBTITLE_FONT = Font(size=11, color="455A64")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill("solid", fgColor="2E7D32")
TOTAL_FONT = Font(bold=True, color="FFFFFF", size=11)
TOTAL_FILL = PatternFill("solid", fgColor="1B5E20")
TOTAL_FILL_GOLD = PatternFill("solid", fgColor="F9A825")
TOTAL_FONT_DARK = Font(bold=True, color="1B5E20", size=11)
ALT_FILL = PatternFill("solid", fgColor="F1F8E9")
THIN_BORDER = Border(
    left=Side(style="thin", color="C8E6C9"),
    right=Side(style="thin", color="C8E6C9"),
    top=Side(style="thin", color="C8E6C9"),
    bottom=Side(style="thin", color="C8E6C9"),
)


def _fetch_all(conn, query: str, params: tuple = ()) -> list[dict[str, Any]]:
    with conn.cursor() as cur:
        cur.execute(query, params)
        return [dict(r) for r in cur.fetchall()]


def _stream_rows(conn, query: str, params: tuple = ()) -> Iterator[dict[str, Any]]:
    with conn.cursor() as cur:
        cur.execute(query, params)
        while True:
            batch = cur.fetchmany(FETCH_BATCH)
            if not batch:
                break
            for row in batch:
                yield dict(row)


def _parse_filters(filters: dict[str, Any], prefix: str = "") -> tuple[list[str], list[Any]]:
    col = f"{prefix}." if prefix else ""
    where = ["1=1"]
    params: list[Any] = []

    status = (filters.get("status") or "").strip()
    if status:
        where.append(f"{col}status = %s")
        params.append(status)

    search = (filters.get("search") or "").strip()
    if search:
        where.append(
            f"({col}order_number ILIKE %s OR {col}shipping_name ILIKE %s OR {col}shipping_phone ILIKE %s)"
        )
        p = f"%{search}%"
        params.extend([p, p, p])

    created_from = filters.get("created_from")
    if created_from:
        where.append(f"{col}created_at >= %s")
        params.append(created_from)

    created_to = filters.get("created_to")
    if created_to:
        try:
            end = datetime.fromisoformat(str(created_to).replace("Z", "+00:00"))
            if end.tzinfo is None:
                end = end.replace(tzinfo=timezone.utc)
            end = end + timedelta(days=1)
            where.append(f"{col}created_at < %s")
            params.append(end)
        except ValueError:
            where.append(f"{col}created_at < %s::date + interval '1 day'")
            params.append(created_to)

    return where, params


def _filter_label(filters: dict[str, Any]) -> str:
    parts: list[str] = []
    if filters.get("status"):
        parts.append(f"Trạng thái: {STATUS_LABELS.get(filters['status'], filters['status'])}")
    if filters.get("search"):
        parts.append(f"Tìm kiếm: {filters['search']}")
    if filters.get("created_from") or filters.get("created_to"):
        parts.append(
            f"Từ {filters.get('created_from') or '...'} đến {filters.get('created_to') or '...'}"
        )
    return " | ".join(parts) if parts else "Tất cả đơn hàng"


def _fetch_summary(
    conn,
    where_sql: str,
    bind: tuple,
    order_where_sql: str,
    order_bind: tuple,
) -> dict[str, Any]:
    with conn.cursor() as cur:
        cur.execute(
            f"""
            SELECT COUNT(*)::int AS total_orders,
                   COALESCE(SUM(CASE WHEN status != 'cancelled' THEN total_amount ELSE 0 END), 0)::float
                       AS total_revenue
            FROM orders
            WHERE {where_sql}
            """,
            bind,
        )
        summary = dict(cur.fetchone() or {})
        cur.execute(
            f"""
            SELECT COALESCE(SUM(oi.quantity), 0)::bigint AS total_items,
                   COUNT(oi.id)::bigint AS item_rows
            FROM order_items oi
            JOIN orders o ON o.id = oi.order_id
            WHERE {order_where_sql}
            """,
            order_bind,
        )
        items = dict(cur.fetchone() or {})
    summary["total_items"] = int(items.get("total_items") or 0)
    summary["item_rows"] = int(items.get("item_rows") or 0)
    return summary


def _write_sheet_title(ws, title: str, subtitle: str | None, ncols: int) -> int:
    ws.cell(row=1, column=1, value=title).font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.row_dimensions[1].height = 30
    if subtitle:
        cell = ws.cell(row=2, column=1, value=subtitle)
        cell.font = SUBTITLE_FONT
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    return 3 if subtitle else 2


def _write_headers(ws, row: int, headers: list[str], widths: list[float]) -> None:
    for col, (text, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[row].height = 24


def _style_data_rows(
    ws,
    data_start: int,
    data_end: int,
    ncols: int,
    col_formats: dict[int, str],
    wrap_cols: set[int] | None = None,
) -> None:
    wrap_cols = wrap_cols or set()
    for row in range(data_start, data_end + 1):
        for col in range(1, ncols + 1):
            cell = ws.cell(row=row, column=col)
            if col in col_formats:
                cell.number_format = col_formats[col]
            cell.border = THIN_BORDER
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=col in wrap_cols,
                horizontal=(
                    "right"
                    if col_formats.get(col) in (VND_FMT, INT_FMT)
                    else "left"
                ),
            )
            if (row - data_start) % 2 == 1:
                cell.fill = ALT_FILL
        ws.row_dimensions[row].height = 20 if not wrap_cols else 22


def _write_total_row(
    ws,
    row: int,
    ncols: int,
    data_start: int,
    data_end: int,
    money_cols: list[int],
    qty_cols: list[int] | None = None,
    label: str = "TỔNG CỘNG",
    fill: PatternFill = TOTAL_FILL,
    font: Font = TOTAL_FONT,
) -> None:
    qty_cols = qty_cols or []
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = font
        cell.fill = fill
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="right" if col in money_cols or col in qty_cols else "left", vertical="center")
        if col == 1:
            cell.value = label
            cell.alignment = Alignment(horizontal="left", vertical="center")
        elif col in money_cols:
            letter = get_column_letter(col)
            cell.value = f"=SUM({letter}{data_start}:{letter}{data_end})"
            cell.number_format = VND_FMT
        elif col in qty_cols:
            letter = get_column_letter(col)
            cell.value = f"=SUM({letter}{data_start}:{letter}{data_end})"
            cell.number_format = INT_FMT


def _apply_filter_freeze(ws, header_row: int, last_col: int, data_end: int) -> None:
    last_letter = get_column_letter(last_col)
    ws.auto_filter.ref = f"A{header_row}:{last_letter}{data_end}"
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)


def _style_overview_sheet(ws, summary: dict[str, Any], filters: dict[str, Any], now: str) -> None:
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 24
    ws.cell(row=1, column=1, value="BÁO CÁO ĐƠN HÀNG — SHOP CAO VĂN SƠN").font = TITLE_FONT
    ws.merge_cells("A1:B1")
    ws.cell(row=2, column=1, value=f"Xuất lúc: {now}").font = SUBTITLE_FONT
    ws.cell(row=3, column=1, value=f"Bộ lọc: {_filter_label(filters)}").font = SUBTITLE_FONT
    metrics = [
        ("Chỉ số", "Giá trị"),
        ("Tổng đơn hàng", summary["total_orders"]),
        ("Tổng doanh thu (không hủy)", summary["total_revenue"]),
        ("Tổng sản phẩm bán", summary["total_items"]),
        ("Số dòng chi tiết SP", summary["item_rows"]),
    ]
    for i, (label, value) in enumerate(metrics, start=5):
        ws.cell(row=i, column=1, value=label)
        val_cell = ws.cell(row=i, column=2, value=value)
        if i == 5:
            for c in (1, 2):
                cell = ws.cell(row=i, column=c)
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL
                cell.border = THIN_BORDER
                cell.alignment = Alignment(horizontal="center")
        else:
            ws.cell(row=i, column=1).border = THIN_BORDER
            val_cell.border = THIN_BORDER
            if label.startswith("Tổng doanh thu"):
                val_cell.number_format = VND_FMT
            elif isinstance(value, int):
                val_cell.number_format = INT_FMT


def build_orders_excel(filters: dict[str, Any], on_progress: ProgressCallback | None = None) -> bytes:
    def report(pct: int, message: str) -> None:
        if on_progress:
            on_progress(pct, message)

    where, params = _parse_filters(filters)
    where_sql = " AND ".join(where)
    order_where, order_params = _parse_filters(filters, prefix="o")
    order_where_sql = " AND ".join(order_where)
    bind = tuple(params)
    order_bind = tuple(order_params)
    now = datetime.now().strftime("%d/%m/%Y %H:%M")

    wb = Workbook()
    wb.remove(wb.active)

    with order_db() as conn:
        report(8, "Đang tổng hợp thống kê...")
        summary = _fetch_summary(conn, where_sql, bind, order_where_sql, order_bind)
        total_orders = int(summary.get("total_orders") or 0)
        item_rows = int(summary.get("item_rows") or 0)
        by_status = _fetch_all(
            conn,
            f"""
            SELECT status, COUNT(*)::int AS count,
                   COALESCE(SUM(total_amount), 0)::float AS revenue
            FROM orders WHERE {where_sql}
            GROUP BY status ORDER BY count DESC
            """,
            bind,
        )
        daily = _fetch_all(
            conn,
            f"""
            SELECT DATE(created_at AT TIME ZONE 'UTC') AS day,
                   COUNT(*)::int AS orders,
                   COALESCE(SUM(CASE WHEN status != 'cancelled' THEN total_amount ELSE 0 END), 0)::float
                       AS revenue
            FROM orders WHERE {where_sql}
            GROUP BY 1 ORDER BY 1 DESC
            """,
            bind,
        )

        report(15, "Đang tạo sheet tổng quan...")
        ws = wb.create_sheet("Tong_quan")
        _style_overview_sheet(ws, summary, filters, now)

        order_headers = [
            "Mã đơn", "Ngày đặt", "Trạng thái", "Khách hàng", "Số điện thoại",
            "Địa chỉ giao", "Tạm tính", "Giảm giá", "Mã coupon", "Tổng thanh toán",
            "User ID", "Cập nhật",
        ]
        order_widths = [16, 20, 14, 24, 16, 40, 18, 16, 14, 20, 38, 20]
        order_ncols = len(order_headers)

        ws_o = wb.create_sheet("Don_hang")
        order_hdr_row = _write_sheet_title(ws_o, "DANH SÁCH ĐƠN HÀNG", None, order_ncols)
        _write_headers(ws_o, order_hdr_row, order_headers, order_widths)
        order_data_start = order_hdr_row + 1

        orders_sql = f"""
            SELECT order_number, id, status, shipping_name, shipping_phone, shipping_address,
                   subtotal_amount, discount_amount, coupon_code, total_amount,
                   user_id, created_at, updated_at
            FROM orders
            WHERE {where_sql}
            ORDER BY created_at DESC
        """
        report(18, "Đang ghi danh sách đơn hàng...")
        order_count = 0
        for o in _stream_rows(conn, orders_sql, bind):
            order_count += 1
            if order_count % 200 == 0 or order_count == total_orders:
                if total_orders > 0:
                    pct = 18 + int(32 * order_count / total_orders)
                    report(pct, f"Đang ghi đơn hàng ({order_count:,}/{total_orders:,})...")
            ws_o.append([
                o.get("order_number") or str(o.get("id", ""))[:8],
                o.get("created_at"),
                STATUS_LABELS.get(o.get("status", ""), o.get("status")),
                o.get("shipping_name"),
                o.get("shipping_phone"),
                o.get("shipping_address"),
                float(o.get("subtotal_amount") or 0),
                float(o.get("discount_amount") or 0),
                o.get("coupon_code") or "",
                float(o.get("total_amount") or 0),
                str(o.get("user_id", "")),
                o.get("updated_at"),
            ])
        if total_orders == 0:
            report(50, "Không có đơn hàng trong bộ lọc")

        order_data_end = ws_o.max_row
        if order_data_end >= order_data_start:
            order_total_row = order_data_end + 1
            _write_total_row(
                ws_o, order_total_row, order_ncols, order_data_start, order_data_end,
                money_cols=[7, 8, 10],
            )
            _style_data_rows(
                ws_o, order_data_start, order_data_end, order_ncols,
                col_formats={2: DATETIME_FMT, 7: VND_FMT, 8: VND_FMT, 10: VND_FMT, 12: DATETIME_FMT},
                wrap_cols={6},
            )
            _apply_filter_freeze(ws_o, order_hdr_row, order_ncols, order_data_end)

        item_headers = [
            "Mã đơn", "Ngày đặt", "Trạng thái đơn", "Tên sản phẩm", "Product ID",
            "Số lượng", "Đơn giá", "Thành tiền",
        ]
        item_widths = [16, 20, 16, 36, 38, 12, 18, 20]
        item_ncols = len(item_headers)

        ws_i = wb.create_sheet("Chi_tiet_SP")
        item_hdr_row = _write_sheet_title(ws_i, "CHI TIẾT SẢN PHẨM THEO ĐƠN", None, item_ncols)
        _write_headers(ws_i, item_hdr_row, item_headers, item_widths)
        item_data_start = item_hdr_row + 1

        items_sql = f"""
            SELECT o.order_number, o.status, o.created_at,
                   oi.product_name_snapshot, oi.product_id, oi.quantity,
                   oi.unit_price, (oi.unit_price * oi.quantity) AS line_total
            FROM order_items oi
            JOIN orders o ON o.id = oi.order_id
            WHERE {order_where_sql}
            ORDER BY o.created_at DESC, oi.product_name_snapshot
        """
        report(52, "Đang ghi chi tiết sản phẩm...")
        item_count = 0
        for it in _stream_rows(conn, items_sql, order_bind):
            item_count += 1
            if item_count % 500 == 0 or item_count == item_rows:
                if item_rows > 0:
                    pct = 52 + int(28 * item_count / item_rows)
                    report(pct, f"Đang ghi chi tiết SP ({item_count:,}/{item_rows:,})...")
            ws_i.append([
                it.get("order_number"),
                it.get("created_at"),
                STATUS_LABELS.get(it.get("status", ""), it.get("status")),
                it.get("product_name_snapshot"),
                str(it.get("product_id", "")),
                int(it.get("quantity") or 0),
                float(it.get("unit_price") or 0),
                float(it.get("line_total") or 0),
            ])
        if item_rows == 0:
            report(80, "Không có chi tiết sản phẩm")

        item_data_end = ws_i.max_row
        if item_data_end >= item_data_start:
            item_total_row = item_data_end + 1
            _write_total_row(
                ws_i, item_total_row, item_ncols, item_data_start, item_data_end,
                money_cols=[8], qty_cols=[6],
            )
            _style_data_rows(
                ws_i, item_data_start, item_data_end, item_ncols,
                col_formats={2: DATETIME_FMT, 6: INT_FMT, 7: VND_FMT, 8: VND_FMT},
                wrap_cols={4},
            )
            _apply_filter_freeze(ws_i, item_hdr_row, item_ncols, item_data_end)

        report(82, "Đang ghi thống kê trạng thái và doanh thu...")
        status_headers = ["Trạng thái", "Nhãn", "Số đơn", "Doanh thu"]
        status_widths = [14, 18, 12, 20]
        status_ncols = len(status_headers)

        ws_s = wb.create_sheet("Trang_thai")
        status_hdr_row = _write_sheet_title(ws_s, "THỐNG KÊ THEO TRẠNG THÁI", None, status_ncols)
        _write_headers(ws_s, status_hdr_row, status_headers, status_widths)
        status_data_start = status_hdr_row + 1
        for row in by_status:
            ws_s.append([
                row.get("status"),
                STATUS_LABELS.get(row.get("status", ""), row.get("status")),
                int(row.get("count") or 0),
                float(row.get("revenue") or 0),
            ])
        status_data_end = ws_s.max_row
        if status_data_end >= status_data_start:
            status_total_row = status_data_end + 1
            _write_total_row(
                ws_s, status_total_row, status_ncols, status_data_start, status_data_end,
                money_cols=[4], qty_cols=[3],
            )
            _style_data_rows(
                ws_s, status_data_start, status_data_end, status_ncols,
                col_formats={3: INT_FMT, 4: VND_FMT},
            )
            _apply_filter_freeze(ws_s, status_hdr_row, status_ncols, status_data_end)

        daily_headers = ["Ngày", "Số đơn", "Doanh thu"]
        daily_widths = [16, 14, 22]
        daily_ncols = len(daily_headers)

        ws_d = wb.create_sheet("Doanh_thu_ngay")
        daily_hdr_row = _write_sheet_title(ws_d, "DOANH THU THEO NGÀY", None, daily_ncols)
        _write_headers(ws_d, daily_hdr_row, daily_headers, daily_widths)
        daily_data_start = daily_hdr_row + 1
        for row in daily:
            ws_d.append([
                row.get("day"),
                int(row.get("orders") or 0),
                float(row.get("revenue") or 0),
            ])
        daily_data_end = ws_d.max_row
        if daily_data_end >= daily_data_start:
            daily_total_row = daily_data_end + 1
            _write_total_row(
                ws_d, daily_total_row, daily_ncols, daily_data_start, daily_data_end,
                money_cols=[3], qty_cols=[2],
                fill=TOTAL_FILL_GOLD,
                font=TOTAL_FONT_DARK,
            )
            _style_data_rows(
                ws_d, daily_data_start, daily_data_end, daily_ncols,
                col_formats={1: DATE_FMT, 2: INT_FMT, 3: VND_FMT},
            )
            _apply_filter_freeze(ws_d, daily_hdr_row, daily_ncols, daily_data_end)

    report(90, "Đang lưu file Excel...")
    buf = io.BytesIO()
    wb.save(buf)
    report(98, "Đang hoàn tất...")
    return buf.getvalue()
