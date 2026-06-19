import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from services import reports


def build_excel_report(period: str) -> bytes:
    overview = reports.overview(period)
    revenue = reports.revenue_chart(period)
    users = reports.users_chart(period)
    reviews = reports.reviews_chart(period)
    top = reports.top_products(period, 20)

    wb = Workbook()
    ws = wb.active
    ws.title = "Tong quan"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1565C0")

    ws["A1"] = "BÁO CÁO SHOP CAO VĂN SƠN"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Kỳ: {period.upper()} | Xuất lúc: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws.merge_cells("A1:E1")

    rows = [
        ("Chỉ số", "Giá trị"),
        ("Doanh thu", overview["orders"]["revenue"]),
        ("Đơn hàng", overview["orders"]["total"]),
        ("Khách mới", overview["users"]["new_customers"]),
        ("Tổng khách hàng", overview["users"]["total_customers"]),
        ("Sản phẩm đang bán", overview["products"]["total_products"]),
        ("Review mới", overview["products"]["new_reviews"]),
        ("Điểm review TB", overview["products"]["avg_rating"]),
    ]
    for i, (k, v) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v)

    _sheet_table(wb.create_sheet("Doanh thu"), "Doanh thu theo thời gian", ["Thời gian", "Doanh thu", "Đơn hàng"],
                 [(r["label"], r["revenue"], r["orders"]) for r in revenue], header_font, header_fill)

    _sheet_table(wb.create_sheet("Nguoi dung"), "Khách đăng ký mới", ["Thời gian", "Số lượng"],
                 [(r["label"], r["count"]) for r in users], header_font, header_fill)

    _sheet_table(wb.create_sheet("Reviews"), "Review & đánh giá", ["Thời gian", "Số review", "Điểm TB"],
                 [(r["label"], r["count"], r["avg_rating"]) for r in reviews], header_font, header_fill)

    _sheet_table(wb.create_sheet("Top SP"), "Sản phẩm bán chạy", ["Tên SP", "SL bán", "Doanh thu", "Số đơn"],
                 [(r["name"], r["quantity_sold"], r["revenue"], r["order_count"]) for r in top], header_font, header_fill)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet_table(ws, title, headers, rows, header_font, header_fill):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=12)
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for r_idx, row in enumerate(rows, start=4):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18
